"""Отчёты: Excel-файл с графиком и выгрузка в Google Таблицу через API (сервисный аккаунт)."""
from __future__ import annotations

import json
import logging
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import httpx
import jwt
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

log = logging.getLogger(__name__)
HEADERS = ["Товар", "Цена", "Было", "Изменение, %", "Мин. за 30 дней", "Цель", "Наличие", "Обновлено", "Ссылка"]


@dataclass(frozen=True)
class ReportRow:
    name: str
    price: float | None
    previous: float | None
    min_30d: float | None
    target: float | None
    in_stock: bool | None
    updated: datetime | None
    url: str
    error: str | None = None

    @property
    def change_pct(self) -> float | None:
        if self.price is None or not self.previous:
            return None
        return round((self.price - self.previous) / self.previous * 100, 1)

    def as_list(self, tz: ZoneInfo) -> list:
        stock = {True: "в наличии", False: "нет", None: "—"}[self.in_stock]
        return [self.name, self.price, self.previous, self.change_pct, self.min_30d, self.target,
                self.error or stock, self.updated.astimezone(tz).strftime("%d.%m.%Y %H:%M") if self.updated else "—", self.url]


def write_excel(rows: list[ReportRow], history: dict[str, list[tuple[datetime, float]]], path: Path,
                tz: ZoneInfo = ZoneInfo("Europe/Moscow")) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Цены"
    ws.append(HEADERS)
    for r in rows:
        ws.append(r.as_list(tz))
    head_fill = PatternFill("solid", fgColor="111827")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 34
    widths = [42, 12, 12, 14, 16, 12, 14, 18, 50]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=6):
        for c in row:
            c.number_format = "#,##0" if c.column != 4 else "+0.0;-0.0;0"
    last = ws.max_row
    if last >= 2:
        ws.conditional_formatting.add(f"D2:D{last}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="15803D", bold=True)))
        ws.conditional_formatting.add(f"D2:D{last}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="B91C1C", bold=True)))
    for i, r in enumerate(rows, start=2):
        ws.cell(i, 9).hyperlink = r.url
        ws.cell(i, 9).style = "Hyperlink"
    ws.freeze_panes = "A2"
    _fit_to_page(ws)
    ws.auto_filter.ref = f"A1:I{max(last, 1)}"

    # История: даты по строкам, товары по столбцам + график
    hs = wb.create_sheet("История")
    names = [r.name for r in rows if history.get(r.url)]
    dates = sorted({d.astimezone(tz).date() for r in rows for d, _ in history.get(r.url, [])})
    hs.append(["Дата", *names])
    for d in dates:
        line: list = [d]
        for r in rows:
            if not history.get(r.url):
                continue
            day_prices = [p for ts, p in history[r.url] if ts.astimezone(tz).date() == d]
            line.append(day_prices[-1] if day_prices else None)
        hs.append(line)
    hs.column_dimensions["A"].width = 12
    for col in range(2, len(names) + 2):
        hs.column_dimensions[get_column_letter(col)].width = 22
    for cell in hs[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    hs.row_dimensions[1].height = 34
    hs.freeze_panes = "B2"
    if dates and names:
        for row in hs.iter_rows(min_row=2, max_col=len(names) + 1):
            row[0].number_format = "DD.MM"
            for c in row[1:]:
                c.number_format = "#,##0"
        chart = LineChart()
        chart.title = "Динамика цен"
        chart.height, chart.width = 9, 22
        chart.y_axis.title = "Цена"
        data = Reference(hs, min_col=2, max_col=len(names) + 1, min_row=1, max_row=len(dates) + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(Reference(hs, min_col=1, min_row=2, max_row=len(dates) + 1))
        hs.add_chart(chart, f"{get_column_letter(len(names) + 3)}2")
    _fit_to_page(hs)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    log.info("Отчёт сохранён: %s", path)
    return path


def _fit_to_page(ws) -> None:
    """Печать на одну страницу в ширину, альбомная ориентация."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


class GoogleSheetsClient:
    """Минимальный клиент Google Sheets API v4 на httpx: авторизация сервисным аккаунтом (JWT)."""

    TOKEN_URL = "https://oauth2.googleapis.com/token"
    SCOPE = "https://www.googleapis.com/auth/spreadsheets"

    def __init__(self, credentials_file: Path, transport: httpx.AsyncBaseTransport | None = None) -> None:
        try:
            info = json.loads(Path(credentials_file).read_text(encoding="utf-8"))
            self._email: str = info["client_email"]
            self._key: str = info["private_key"]
        except (OSError, KeyError, json.JSONDecodeError) as exc:
            raise ValueError(f"Некорректный файл ключа сервисного аккаунта: {credentials_file}") from exc
        self._client = httpx.AsyncClient(timeout=20, transport=transport)
        self._token: str | None = None
        self._token_exp = 0.0

    async def _access_token(self) -> str:
        if self._token and time.time() < self._token_exp - 60:
            return self._token
        now = int(time.time())
        assertion = jwt.encode({"iss": self._email, "scope": self.SCOPE, "aud": self.TOKEN_URL, "iat": now, "exp": now + 3600},
                               self._key, algorithm="RS256")
        resp = await self._client.post(self.TOKEN_URL, data={"grant_type": "urn:ietf:params:oauth:grant-type:jwt-bearer",
                                                             "assertion": assertion})
        resp.raise_for_status()
        body = resp.json()
        self._token, self._token_exp = body["access_token"], now + int(body.get("expires_in", 3600))
        return self._token

    async def write(self, sheet_id: str, rows: list[list], sheet_name: str = "Цены") -> None:
        token = await self._access_token()
        base = f"https://sheets.googleapis.com/v4/spreadsheets/{sheet_id}/values"
        headers = {"Authorization": f"Bearer {token}"}
        rng = f"'{sheet_name}'!A1"
        clear = await self._client.post(f"{base}/'{sheet_name}':clear", headers=headers)
        clear.raise_for_status()
        resp = await self._client.put(f"{base}/{rng}", params={"valueInputOption": "USER_ENTERED"},
                                      headers=headers, json={"range": rng, "majorDimension": "ROWS", "values": rows})
        resp.raise_for_status()
        log.info("Google Таблица обновлена: %d строк", len(rows))

    async def close(self) -> None:
        await self._client.aclose()
