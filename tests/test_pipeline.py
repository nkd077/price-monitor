"""Загрузчик, уведомления, история и отчёты — без реальной сети (httpx.MockTransport)."""
import json
import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path

import httpx
from cryptography.hazmat.primitives import serialization
from cryptography.hazmat.primitives.asymmetric import rsa
from openpyxl import load_workbook

from price_monitor.alerts import TelegramNotifier, evaluate
from price_monitor.config import Product, Selectors, Settings
from price_monitor.export import GoogleSheetsClient, write_excel
from price_monitor.fetcher import FetchError, Fetcher
from price_monitor.monitor import run_once
from price_monitor.storage import Observation, Storage

FIX = Path(__file__).parent / "fixtures"


async def no_sleep(_: float) -> None:
    return None


def settings(products) -> Settings:
    return Settings(telegram_token=None, telegram_chat_id=None, drop_alert_percent=5, concurrency=3, per_host_delay=0,
                    timeout=5, max_retries=2, respect_robots=True, db_path=Path(":memory:"), google_sheet_id=None,
                    google_credentials=None, products=tuple(products))


class FetcherTest(unittest.IsolatedAsyncioTestCase):
    async def test_retries_then_succeeds(self) -> None:
        calls = {"n": 0}

        def handler(req: httpx.Request) -> httpx.Response:
            if req.url.path == "/robots.txt":
                return httpx.Response(404)
            calls["n"] += 1
            return httpx.Response(503) if calls["n"] < 3 else httpx.Response(200, text="ok")

        async with Fetcher(transport=httpx.MockTransport(handler), sleep=no_sleep, per_host_delay=0, max_retries=3) as f:
            self.assertEqual(await f.get("https://a.test/item"), "ok")
        self.assertEqual(calls["n"], 3)

    async def test_404_not_retried(self) -> None:
        calls = {"n": 0}

        def handler(req: httpx.Request) -> httpx.Response:
            if req.url.path == "/robots.txt":
                return httpx.Response(404)
            calls["n"] += 1
            return httpx.Response(404)

        async with Fetcher(transport=httpx.MockTransport(handler), sleep=no_sleep, per_host_delay=0) as f:
            with self.assertRaises(FetchError) as ctx:
                await f.get("https://a.test/missing")
        self.assertIn("HTTP 404", str(ctx.exception))
        self.assertEqual(calls["n"], 1)

    async def test_robots_disallow_respected(self) -> None:
        def handler(req: httpx.Request) -> httpx.Response:
            if req.url.path == "/robots.txt":
                return httpx.Response(200, text="User-agent: *\nDisallow: /private/\n")
            return httpx.Response(200, text="page")

        async with Fetcher(transport=httpx.MockTransport(handler), sleep=no_sleep, per_host_delay=0) as f:
            self.assertEqual(await f.get("https://a.test/public/1"), "page")
            with self.assertRaises(FetchError) as ctx:
                await f.get("https://a.test/private/1")
        self.assertIn("robots.txt", str(ctx.exception))

    async def test_rotates_user_agent(self) -> None:
        agents = set()

        def handler(req: httpx.Request) -> httpx.Response:
            agents.add(req.headers["User-Agent"])
            return httpx.Response(200, text="x")

        async with Fetcher(transport=httpx.MockTransport(handler), sleep=no_sleep, per_host_delay=0, respect_robots=False) as f:
            for i in range(40):
                await f.get(f"https://a.test/{i}")
        self.assertGreater(len(agents), 1)


class AlertsTest(unittest.TestCase):
    def obs(self, price: float, stock=True) -> Observation:
        return Observation("u", "Товар", None, price, "RUB", stock, datetime.now(timezone.utc))

    def test_rules(self) -> None:
        p = Product("Товар", "https://x.test/p", target_price=4500)
        self.assertEqual([a.kind for a in evaluate(p, self.obs(5000), self.obs(4800), 5)], [])            # −4%
        self.assertEqual([a.kind for a in evaluate(p, self.obs(5000), self.obs(4700), 5)], ["drop"])      # −6%
        self.assertEqual([a.kind for a in evaluate(p, self.obs(4700), self.obs(4400), 50)], ["target"])
        self.assertEqual([a.kind for a in evaluate(p, self.obs(4400), self.obs(4300), 50)], [])           # цель уже была
        self.assertEqual([a.kind for a in evaluate(p, self.obs(5000, False), self.obs(5000, True), 5)], ["back_in_stock"])

    def test_render_formats_money(self) -> None:
        p = Product("Кроссовки, 42 размер", "https://x.test/p")
        text = evaluate(p, self.obs(12990), self.obs(9990), 5)[0].render()
        self.assertIn("12 990 → <b>9 990 ₽</b>", text)
        self.assertIn("Кроссовки, 42 размер", text)


class PipelineTest(unittest.IsolatedAsyncioTestCase):
    async def test_two_runs_price_drop_alert_and_report(self) -> None:
        pages = {"/runner": (FIX / "jsonld_shop.html").read_text(encoding="utf-8"),
                 "/lamp": (FIX / "plain_shop.html").read_text(encoding="utf-8")}

        def handler(req: httpx.Request) -> httpx.Response:
            if req.url.path == "/robots.txt":
                return httpx.Response(404)
            if req.url.path == "/gone":
                return httpx.Response(404)
            return httpx.Response(200, text=pages[req.url.path])

        products = [Product("Кроссовки", "https://shop.test/runner", target_price=4500),
                    Product("Лампа", "https://lamps.test/lamp", selectors=Selectors(price=".price__current", in_stock=".buy-button")),
                    Product("Снят с продажи", "https://shop.test/gone")]
        s = settings(products)
        storage = Storage(":memory:")
        sent: list[str] = []

        class Notifier:
            async def send(self, text: str) -> bool:
                sent.append(text)
                return True

        t0 = datetime(2030, 1, 7, 9, tzinfo=timezone.utc)
        async with Fetcher(transport=httpx.MockTransport(handler), sleep=no_sleep, per_host_delay=0) as f:
            r1 = await run_once(s, f, storage, Notifier(), now=t0)
            self.assertEqual(r1.errors, 1)
            self.assertEqual(sent, [])
            pages["/runner"] = pages["/runner"].replace('"4990.00"', '"4390.00"')
            r2 = await run_once(s, f, storage, Notifier(), now=t0 + timedelta(days=1))
        kinds = sorted(a.kind for a in r2.alerts)
        self.assertEqual(kinds, ["drop", "target"])
        self.assertEqual(len(sent), 2)
        runner = next(r for r in r2.rows if r.name == "Кроссовки")
        self.assertEqual((runner.price, runner.previous, runner.min_30d, runner.change_pct), (4390.0, 4990.0, 4390.0, -12.0))
        gone = next(r for r in r2.rows if r.name == "Снят с продажи")
        self.assertEqual(gone.error, "HTTP 404")

        history = {p.url: [(o.seen_at, o.price) for o in await storage.history(p.url, now=t0 + timedelta(days=1))] for p in products}
        with tempfile.TemporaryDirectory() as tmp:
            path = write_excel(r2.rows, history, Path(tmp) / "report.xlsx")
            wb = load_workbook(path)
            ws = wb["Цены"]
            self.assertEqual(ws["A2"].value, "Кроссовки")
            self.assertEqual(ws["B2"].value, 4390.0)
            self.assertEqual(ws["D2"].value, -12.0)
            self.assertEqual(ws["G4"].value, "HTTP 404")
            self.assertEqual(len(wb["История"]._charts), 1)
        storage.close()


class TelegramAndSheetsTest(unittest.IsolatedAsyncioTestCase):
    async def test_telegram_payload(self) -> None:
        seen = {}

        def handler(req: httpx.Request) -> httpx.Response:
            seen["url"] = str(req.url)
            seen["body"] = json.loads(req.content)
            return httpx.Response(200, json={"ok": True})

        n = TelegramNotifier("123:ABC", "-100500", transport=httpx.MockTransport(handler))
        self.assertTrue(await n.send("<b>тест</b>"))
        await n.close()
        self.assertTrue(seen["url"].endswith("/bot123:ABC/sendMessage"))
        self.assertEqual(seen["body"]["chat_id"], "-100500")
        self.assertEqual(seen["body"]["parse_mode"], "HTML")

    async def test_sheets_auth_and_write(self) -> None:
        key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
        pem = key.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8, serialization.NoEncryption()).decode()
        calls = []

        def handler(req: httpx.Request) -> httpx.Response:
            calls.append((req.method, req.url.path, req.headers.get("Authorization")))
            if req.url.host == "oauth2.googleapis.com":
                self.assertIn(b"jwt-bearer", req.content)
                return httpx.Response(200, json={"access_token": "tok", "expires_in": 3600})
            if req.method == "PUT":
                body = json.loads(req.content)
                self.assertEqual(body["values"][1][0], "Кроссовки")
            return httpx.Response(200, json={})

        with tempfile.TemporaryDirectory() as tmp:
            cred = Path(tmp) / "sa.json"
            cred.write_text(json.dumps({"client_email": "bot@proj.iam.gserviceaccount.com", "private_key": pem}))
            client = GoogleSheetsClient(cred, transport=httpx.MockTransport(handler))
            await client.write("SHEET", [["Товар", "Цена"], ["Кроссовки", 4390]])
            await client.close()
        self.assertEqual(calls[0][1], "/token")
        self.assertTrue(all(c[2] == "Bearer tok" for c in calls[1:]))
        self.assertEqual([c[0] for c in calls[1:]], ["POST", "PUT"])


if __name__ == "__main__":
    unittest.main()
